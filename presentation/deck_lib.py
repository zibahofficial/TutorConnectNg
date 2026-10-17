"""
TutorConnect Nigeria — Edumentry AI Hackathon 1.0
Shared drawing helpers for the PowerPoint deck builder.

Design system mirrors the live product (app/globals.css + tailwind.config.js):
  navy  #1E3A8A  →  blue #2563EB   (primary gradient)
  emerald #10B981 (positive)      amber #F59E0B (warm accent)
  slate scale for surfaces/text    white cards
"""

from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

# ---------------------------------------------------------------- palette ---
NAVY = RGBColor(0x1E, 0x3A, 0x8A)
NAVY_DK = RGBColor(0x11, 0x1E, 0x46)
NAVY_800 = RGBColor(0x17, 0x2B, 0x63)
BLUE = RGBColor(0x25, 0x63, 0xEB)
BLUE_LT = RGBColor(0xDB, 0xEA, 0xFE)
EMERALD = RGBColor(0x10, 0xB9, 0x81)
EMERALD_DK = RGBColor(0x05, 0x96, 0x69)
EMERALD_LT = RGBColor(0xD1, 0xFA, 0xE5)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)
AMBER_LT = RGBColor(0xFE, 0xF3, 0xC7)
ROSE = RGBColor(0xE1, 0x1D, 0x48)
ROSE_LT = RGBColor(0xFF, 0xE4, 0xE6)
SLATE_900 = RGBColor(0x0F, 0x17, 0x2A)
SLATE_700 = RGBColor(0x33, 0x41, 0x55)
SLATE_600 = RGBColor(0x47, 0x55, 0x69)
SLATE_500 = RGBColor(0x64, 0x74, 0x8B)
SLATE_400 = RGBColor(0x94, 0xA3, 0xB8)
SLATE_300 = RGBColor(0xCB, 0xD5, 0xE1)
SLATE_200 = RGBColor(0xE2, 0xE8, 0xF0)
SLATE_100 = RGBColor(0xF1, 0xF5, 0xF9)
SLATE_50 = RGBColor(0xF8, 0xFA, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Segoe UI"
FONT_LIGHT = "Segoe UI Light"
FONT_SB = "Segoe UI Semibold"

SW = 13.333          # slide width  (in)
SH = 7.5             # slide height (in)
M = 0.62             # page margin
CW = SW - 2 * M      # content width

# ------------------------------------------------------------- primitives ---


def _no_line(shape):
    shape.line.fill.background()


def solid(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def set_adj(shape, value):
    """Corner radius for rounded rectangles (0..0.5 of the short side)."""
    try:
        shape.adjustments[0] = value
    except (IndexError, ValueError):
        pass


def rect(slide, x, y, w, h, fill=WHITE, line=None, lw=1.0, radius=0.08,
         shape_type=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=False):
    sh = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
        set_adj(sh, radius)
    if fill is None:
        sh.fill.background()
    else:
        solid(sh, fill)
    if line is None:
        _no_line(sh)
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    if not shadow:
        sh.shadow.inherit = False
    sh.text_frame.word_wrap = True
    return sh


def grad_rect(slide, x, y, w, h, c1=NAVY, c2=BLUE, angle=0, radius=0.0,
              shape_type=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
        set_adj(sh, radius)
    sh.fill.gradient()
    stops = sh.fill.gradient_stops
    stops[0].color.rgb = c1
    stops[0].position = 0.0
    stops[1].color.rgb = c2
    stops[1].position = 1.0
    try:
        sh.fill.gradient_angle = angle
    except Exception:
        pass
    _no_line(sh)
    sh.shadow.inherit = False
    return sh


def grad_overlay(slide, x, y, w, h, c1=NAVY_DK, c2=NAVY, angle=0,
                 shape_type=MSO_SHAPE.RECTANGLE):
    sh = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.gradient()
    stops = sh.fill.gradient_stops
    stops[0].color.rgb = c1
    stops[0].position = 0.0
    stops[1].color.rgb = c2
    stops[1].position = 1.0
    try:
        sh.fill.gradient_angle = angle
    except Exception:
        pass
    sh.fill.transparency = 0.0
    _no_line(sh)
    sh.shadow.inherit = False
    return sh


def line(slide, x1, y1, x2, y2, color=SLATE_200, width=1.0, dash=False):
    from pptx.enum.shapes import MSO_CONNECTOR
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                     Inches(x2), Inches(y2))
    cn.line.color.rgb = color
    cn.line.width = Pt(width)
    if dash:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        cn.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return cn


def tb(slide, x, y, w, h, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    tf.paragraphs[0].alignment = align
    return box, tf


def para(tf, text="", size=14, color=SLATE_700, bold=False, align=PP_ALIGN.LEFT,
         space_before=0, space_after=0, line=1.18, font=FONT, italic=False,
         first=False, caps=False, spacing=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line
    if text:
        r = p.add_run()
        r.text = text
        f = r.font
        f.name = font
        f.size = Pt(size)
        f.bold = bold
        f.italic = italic
        f.color.rgb = color
        if caps:
            r.font._rPr.set('cap', 'all')
        if spacing is not None:
            r.font._rPr.set('spc', str(int(spacing * 100)))
    return p


def run(p, text, size=14, color=SLATE_700, bold=False, font=FONT, italic=False,
        caps=False, spacing=None):
    r = p.add_run()
    r.text = text
    f = r.font
    f.name = font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    if caps:
        r.font._rPr.set('cap', 'all')
    if spacing is not None:
        r.font._rPr.set('spc', str(int(spacing * 100)))
    return r


def text_at(slide, x, y, w, h, text, size=14, color=SLATE_700, bold=False,
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT, italic=False,
            caps=False, spacing=None, line=1.18):
    box, tf = tb(slide, x, y, w, h, align, anchor)
    para(tf, text, size, color, bold, align, font=font, italic=italic,
         first=True, caps=caps, spacing=spacing, line=line)
    return box


# ------------------------------------------------------------ measurement ---

def est_lines(text, width_in, size_pt, bold=False):
    """Rough line count for a run of text in a box of the given width."""
    if not text:
        return 1
    cpl = max(int((width_in * 72) / (size_pt * (0.505 if bold else 0.485))), 6)
    total = 0
    for chunk in text.split("\n"):
        total += max(1, -(-len(chunk) // cpl))
    return total


def text_h(text, width_in, size_pt, line=1.18, bold=False):
    return est_lines(text, width_in, size_pt, bold) * size_pt * line / 72.0


# --------------------------------------------------------------- chrome -----

def new_slide(prs, bg=WHITE):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bgr = rect(s, -0.05, -0.05, SW + 0.1, SH + 0.1, fill=bg,
               shape_type=MSO_SHAPE.RECTANGLE)
    bgr.shadow.inherit = False
    return s


def header(slide, kicker, title, accent=BLUE, title_size=30, sub=None,
           title_color=SLATE_900):
    """Eyebrow + title + optional sub-line, with a short accent rule."""
    rect(slide, M, 0.46, 0.055, 0.19, fill=accent)
    text_at(slide, M + 0.16, 0.42, CW - 0.2, 0.26, kicker.upper(), size=10.5,
            color=accent, bold=True, caps=True, spacing=1.2)
    text_at(slide, M, 0.72, CW, 0.46, title, size=title_size, color=title_color,
            bold=True, font=FONT_SB, line=1.0)
    y = 1.32
    if sub:
        text_at(slide, M, y, CW - 0.4, 0.3, sub, size=13.5, color=SLATE_500, line=1.2)
        y += 0.34
    return y


def footer(slide, n, total=16, note="TutorConnect Nigeria  ·  Edumentry AI Hackathon 1.0",
           color=SLATE_400):
    line(slide, M, SH - 0.52, SW - M, SH - 0.52, color=SLATE_200, width=0.75)
    text_at(slide, M, SH - 0.42, CW - 1.4, 0.24, note, size=9.5, color=color)
    text_at(slide, SW - M - 1.3, SH - 0.42, 1.3, 0.24, f"{n} / {total}", size=9.5,
            color=color, align=PP_ALIGN.RIGHT, bold=True)


def presenter_tag(slide, name, role, x=None, y=0.32, w=4.6, dark=False):
    """Small 'Presented by' pill for the top-right of a presenter slide."""
    if len(name) > 25 and w == 4.6:
        w = 5.2
    x = SW - M - w if x is None else x
    bg = RGBColor(0x1B, 0x2E, 0x63) if dark else BLUE_LT
    fg = WHITE if dark else NAVY
    sh = rect(slide, x, y, w, 0.36, fill=bg, radius=0.5)
    tf = sh.text_frame
    tf.margin_left = Inches(0.16)
    tf.margin_right = Inches(0.10)
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run(p, "Presenter  ", size=8.5, color=(SLATE_400 if not dark else SLATE_300),
        bold=True, caps=True, spacing=0.8)
    run(p, name, size=(9.5 if len(name) > 25 else 10.5), color=fg, bold=True)
    return sh


# ----------------------------------------------------------------- cards ----

def card(slide, x, y, w, h, fill=WHITE, border=SLATE_200, radius=0.06, accent=None,
         accent_h=0.055):
    sh = rect(slide, x, y, w, h, fill=fill, line=border, lw=1.0, radius=radius)
    if accent:
        rect(slide, x, y, w, accent_h, fill=accent, radius=0.5)
    return sh


def icon_badge(slide, x, y, size, glyph, bg=BLUE_LT, fg=NAVY, fsize=15, radius=0.28):
    sh = rect(slide, x, y, size, size, fill=bg, radius=radius)
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run(p, glyph, size=fsize, color=fg, bold=True)
    return sh


def bullet_list(slide, x, y, w, items, size=12.5, gap=0.14, marker=BLUE,
                lead_color=SLATE_900, body_color=SLATE_600, bar=0.038):
    """items: list of (lead, body) or plain strings. Accent bar + wrapped text."""
    cy = y
    for it in items:
        if isinstance(it, tuple):
            lead, body = it
        else:
            lead, body = None, it
        full = (lead + " " if lead else "") + (body or "")
        h = text_h(full, w - 0.26, size, 1.2)
        rect(slide, x, cy + 0.045, bar, max(h - 0.05, 0.12), fill=marker, radius=0.5)
        box, tf = tb(slide, x + 0.26, cy, w - 0.26, h + 0.05)
        p = tf.paragraphs[0]
        p.line_spacing = 1.2
        if lead:
            run(p, lead, size=size, color=lead_color, bold=True)
            if body:
                run(p, "  " + body, size=size, color=body_color)
        else:
            run(p, body, size=size, color=body_color)
        cy += h + gap
    return cy


def chip(slide, x, y, w, h, text, fill=BLUE_LT, fg=NAVY, size=10.5, bold=True,
         radius=0.5, align=PP_ALIGN.CENTER):
    sh = rect(slide, x, y, w, h, fill=fill, radius=radius)
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Inches(0.07)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    run(p, text, size=size, color=fg, bold=bold)
    return sh


def arrow(slide, x, y, w, h, color=SLATE_300, shape_type=MSO_SHAPE.RIGHT_ARROW):
    sh = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    solid(sh, color)
    _no_line(sh)
    sh.shadow.inherit = False
    return sh


def chevron(slide, x, y, w, h, text, fill=NAVY, fg=WHITE, size=11.5, bold=True,
            sub=None, sub_color=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    solid(sh, fill)
    _no_line(sh)
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.16)
    tf.margin_right = Inches(0.05)
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.line_spacing = 1.05
    run(p, text, size=size, color=fg, bold=bold)
    if sub:
        para(tf, sub, size=8.5, color=sub_color or fg, align=PP_ALIGN.CENTER,
             line=1.05, space_before=1)
    return sh


def numbered_node(slide, x, y, d, n, fill=NAVY, fg=WHITE, size=13):
    sh = rect(slide, x, y, d, d, fill=fill, radius=0.5, shape_type=MSO_SHAPE.OVAL)
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run(p, str(n), size=size, color=fg, bold=True)
    return sh


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text
